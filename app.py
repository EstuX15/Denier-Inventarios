# -*- coding: utf-8 -*-
import io
import os
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from flask import Flask, render_template_string, request, redirect, url_for, session, send_file, send_from_directory
from datetime import datetime
import openpyxl

app = Flask(__name__)
app.secret_key = 'clave_secreta_denier_unificada_total'

# Configuración de carpeta para archivos subidos (Devoluciones)
UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# ==========================================
# BASE DE DATOS Y MEMORIA UNIFICADA
# ==========================================
usuarios_db = [
    { 'id': 1, 'email': 'controlprod_1@deniersa.com', 'password': '123', 'nombre': 'Administrador General', 'modulos': ['admin', 'hilos', 'cruda', 'devo'] },
    { 'id': 2, 'email': 'bodega.hilos@deniersa.com', 'password': '123', 'nombre': 'Operador Hilos', 'modulos': ['hilos'] },
    { 'id': 3, 'email': 'bodega.cruda@deniersa.com', 'password': '123', 'nombre': 'Operador Tela Cruda', 'modulos': ['cruda'] },
    { 'id': 4, 'email': 'bodega.devoluciones@deniersa.com', 'password': '123', 'nombre': 'Operador Devoluciones', 'modulos': ['devo'] }
]

inventario_hilos = [
    { 'id': 1, 'tipo': 'Denier', 'oc': 'P04150', 'codigo': 'H-COT-20', 'calibre': '20/1', 'composicion': 'Algodon 100%', 'proveedor': 'Hilos del Sur', 'lote': 'L-442', 'kilos': 450.5, 'conos': 120, 'ubicacion': 'A-1' },
    { 'id': 2, 'tipo': 'Denier', 'oc': '4555', 'codigo': 'H-COT-20', 'calibre': '20/1', 'composicion': 'Algodon 100%', 'proveedor': 'Hilos del Sur', 'lote': 'L-442', 'kilos': 35.0, 'conos': 10, 'ubicacion': 'B-2' }
]
packing_history_hilos = []
contador_packing_hilos = 1

inventario_tela = [
    { 'id': 1, 'oc': 'OC-9012', 'proveedor': 'Knitopia', 'lote_tela': 'TC-101', 'rollo': 'R-001', 'kilos': 250.5, 'ubicacion': 'Bodega A-1' },
    { 'id': 2, 'oc': 'OC-9013', 'proveedor': 'Magic Dyeing', 'lote_tela': 'TC-102', 'rollo': 'R-002', 'kilos': 45.0, 'ubicacion': 'Bodega B-2' }
]
packing_tela_history = []
contador_packing_tela = 1

devoluciones_procesada = [
    { 
        'id': 1, 
        'cliente': 'Magic Apparel', 
        'lote': '20452026-8', 
        'tela': 'Interlock, 100% Algodón', 
        'color': 'White, Midnight Navy', 
        'rollos': 3, 
        'kilos': 75.50, 
        'yardas': 240.00, 
        'calidad': 'Rechazada', 
        'ubicacion': 'Rack 3', 
        'archivo': None 
    }
]

# ==========================================
# PLANTILLAS HTML ORIGINALES
# ==========================================
LOGIN_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head><meta charset="UTF-8"><title>DENIER - INVENTARIO</title><script src="https://cdn.tailwindcss.com"></script></head>
<body class="bg-black text-slate-100 min-h-screen flex items-center justify-center p-6">
    <div class="bg-slate-900 border border-slate-800 p-8 rounded-2xl shadow-2xl w-full max-w-sm">
        <div class="text-center mb-6">
            <h1 class="text-2xl font-black tracking-widest text-white uppercase">DENIER</h1>
            <p class="text-xs font-semibold text-slate-400 tracking-wider uppercase mt-1">INVENTARIO UNIFICADO</p>
            <div class="border-b border-slate-800 my-4"></div>
        </div>
        {% if error %}<div class="bg-red-950 border border-red-800 text-red-300 p-3 rounded-lg mb-4 text-xs font-bold text-center">{{ error }}</div>{% endif %}
        <form action="/login" method="POST" class="space-y-4">
            <div>
                <label class="block text-xs font-bold mb-1.5 text-slate-300">Email</label>
                <input type="email" name="email" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-sm text-white" placeholder="controlprod_1@deniersa.com">
            </div>
            <div>
                <label class="block text-xs font-bold mb-1.5 text-slate-300">Password</label>
                <input type="password" name="password" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-sm text-white" placeholder="••••••••">
            </div>
            <button type="submit" class="w-full bg-teal-700 hover:bg-teal-600 text-white py-2.5 rounded-lg font-bold text-sm">Log in</button>
        </form>
        <div class="flex justify-between items-center mt-6 text-xs text-slate-400">
            <span class="text-slate-500">¿Problemas?</span>
            <a href="/reset-password-form" class="text-teal-400 font-semibold hover:underline">Reset Password</a>
        </div>
    </div>
</body>
</html>
"""

HOME_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head><meta charset="UTF-8"><title>DENIER - Panel Principal</title><script src="https://cdn.tailwindcss.com"></script></head>
<body class="bg-black text-slate-100 min-h-screen p-8">
    <div class="max-w-6xl mx-auto">
        <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl mb-8 flex justify-between items-center">
            <div>
                <h1 class="text-2xl font-black tracking-wide text-white">DENIER - INVENTARIO</h1>
                <p class="text-sm text-slate-400">Usuario: <b class="text-teal-400">{{ session.get('nombre') }}</b> ({{ session.get('email') }})</p>
            </div>
            <a href="/logout" class="bg-red-950 text-red-300 border border-red-800 px-4 py-2 rounded-xl text-sm font-semibold">Cerrar Sesión</a>
        </div>
        <div class="grid grid-cols-1 md:grid-cols-3 gap-6 mb-8">
            {% if 'admin' in session.get('modulos') or 'hilos' in session.get('modulos') %}
            <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl flex flex-col justify-between">
                <div>
                    <h2 class="text-xl font-bold text-teal-400 mb-2">Módulo de Hilos</h2>
                    <p class="text-sm text-slate-400 mb-4">Control de inventario Denier y Artitex, packing lists y devoluciones.</p>
                </div>
                <a href="/hilos" class="bg-teal-700 hover:bg-teal-600 text-white text-center py-3 rounded-xl font-medium shadow">Entrar a Hilos</a>
            </div>
            {% endif %}
            {% if 'admin' in session.get('modulos') or 'cruda' in session.get('modulos') %}
            <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl flex flex-col justify-between">
                <div>
                    <h2 class="text-xl font-bold text-teal-400 mb-2">Módulo de Tela Cruda</h2>
                    <p class="text-sm text-slate-400 mb-4">Control de rollos, cargas masivas Knitopia y reingresos Magic Dyeing.</p>
                </div>
                <a href="/cruda" class="bg-teal-700 hover:bg-teal-600 text-white text-center py-3 rounded-xl font-medium shadow">Entrar a Tela Cruda</a>
            </div>
            {% endif %}
            {% if 'admin' in session.get('modulos') or 'devo' in session.get('modulos') %}
            <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl flex flex-col justify-between">
                <div>
                    <h2 class="text-xl font-bold text-rose-400 mb-2">Módulo Devoluciones</h2>
                    <p class="text-sm text-slate-400 mb-4">Control de devoluciones de tela procesada y gestión de archivos PDF.</p>
                </div>
                <a href="/devo" class="bg-rose-700 hover:bg-rose-600 text-white text-center py-3 rounded-xl font-medium shadow">Entrar a Devoluciones</a>
            </div>
            {% endif %}
        </div>
        {% if 'admin' in session.get('modulos') %}
        <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl">
            <h2 class="text-xl font-bold text-purple-400 mb-4">Gestión de Usuarios y Accesos - DENIER</h2>
            <form action="/crear-usuario" method="POST" class="grid grid-cols-1 md:grid-cols-3 gap-4 mb-8 bg-black p-5 rounded-xl border border-slate-800">
                <div>
                    <label class="block text-xs mb-1 font-semibold text-slate-300">Correo Corporativo</label>
                    <input type="email" name="email" required class="w-full p-2.5 rounded-lg border border-slate-700 bg-slate-900 text-sm text-white" placeholder="usuario@deniersa.com">
                </div>
                <div>
                    <label class="block text-xs mb-1 font-semibold text-slate-300">Nombre Completo</label>
                    <input type="text" name="nombre" required class="w-full p-2.5 rounded-lg border border-slate-700 bg-slate-900 text-sm text-white" placeholder="Nombre...">
                </div>
                <div>
                    <label class="block text-xs mb-1 font-semibold text-slate-300">Contraseña Temporal</label>
                    <input type="password" name="password" required class="w-full p-2.5 rounded-lg border border-slate-700 bg-slate-900 text-sm text-white" placeholder="Clave...">
                </div>
                <div class="md:col-span-3">
                    <label class="block text-xs mb-2 font-semibold text-slate-300">Seleccionar Módulos Permitidos:</label>
                    <div class="flex gap-6">
                        <label class="flex items-center space-x-2 text-sm cursor-pointer text-slate-300"><input type="checkbox" name="modulos" value="hilos" class="w-4 h-4 text-teal-600 rounded"><span>Bodega Hilos</span></label>
                        <label class="flex items-center space-x-2 text-sm cursor-pointer text-slate-300"><input type="checkbox" name="modulos" value="cruda" class="w-4 h-4 text-teal-600 rounded"><span>Bodega Tela Cruda</span></label>
                        <label class="flex items-center space-x-2 text-sm cursor-pointer text-slate-300"><input type="checkbox" name="modulos" value="devo" class="w-4 h-4 text-rose-600 rounded"><span>Devoluciones</span></label>
                        <label class="flex items-center space-x-2 text-sm cursor-pointer text-slate-300"><input type="checkbox" name="modulos" value="admin" class="w-4 h-4 text-purple-600 rounded"><span>Administrador Total</span></label>
                    </div>
                </div>
                <div class="md:col-span-3 flex justify-end">
                    <button type="submit" class="bg-purple-700 hover:bg-purple-600 text-white px-6 py-2.5 rounded-xl text-sm font-medium">Registrar Nuevo Usuario</button>
                </div>
            </form>
            <h3 class="font-bold text-sm text-slate-300 mb-3">Lista de Usuarios Activos y Permisos</h3>
            <table class="w-full text-left text-sm">
                <thead>
                    <tr class="border-b border-slate-800 text-slate-400"><th class="p-3">Correo</th><th class="p-3">Nombre</th><th class="p-3">Contraseña</th><th class="p-3">Módulos</th><th class="p-3 text-center">Acciones</th></tr>
                </thead>
                <tbody class="divide-y divide-slate-800">
                    {% for u in usuarios_db %}
                    <tr>
                        <td class="p-3 font-medium text-teal-400">{{ u.email }}</td>
                        <td class="p-3 text-slate-200">{{ u.nombre }}</td>
                        <td class="p-3 text-amber-300 font-mono font-bold">{{ u.password }}</td>
                        <td class="p-3">{% for m in u.modulos %}<span class="bg-slate-800 text-slate-300 border border-slate-700 px-2 py-0.5 rounded text-xs font-bold uppercase mr-1">{{ m }}</span>{% endfor %}</td>
                        <td class="p-3 text-center">
                            <div class="flex justify-center gap-2">
                                <a href="/editar-usuario/{{ u.id }}" class="bg-amber-950 text-amber-300 border border-amber-800 px-3 py-1 rounded-lg text-xs font-semibold">Editar</a>
                                {% if u.email != 'controlprod_1@deniersa.com' %}
                                <a href="/eliminar-usuario/{{ u.id }}" onclick="return confirm('¿Seguro deseas eliminar?');" class="bg-red-950 text-red-300 border border-red-800 px-3 py-1 rounded-lg text-xs font-semibold">Eliminar</a>
                                {% endif %}
                            </div>
                        </td>
                    </tr>
                    {% endfor %}
                </tbody>
            </table>
        </div>
        {% endif %}
    </div>
</body>
</html>
"""

RESET_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head><meta charset="UTF-8"><title>Recuperar Contraseña</title><script src="https://cdn.tailwindcss.com"></script></head>
<body class="bg-black text-slate-100 min-h-screen flex items-center justify-center p-6">
    <div class="bg-slate-900 border border-slate-800 p-8 rounded-2xl shadow-2xl w-full max-w-sm">
        <h1 class="text-xl font-bold text-white mb-2">Restablecer Contraseña</h1>
        {% if error %}<div class="bg-red-950 border border-red-800 text-red-300 p-3 rounded mb-4 text-xs font-bold text-center">{{ error }}</div>{% endif %}
        {% if mensaje %}<div class="bg-emerald-950 border border-emerald-800 text-emerald-300 p-3 rounded mb-4 text-xs font-bold text-center">{{ mensaje }}</div>{% endif %}
        <form action="/procesar-reset" method="POST" class="space-y-4">
            <input type="email" name="email" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-sm text-white" placeholder="Correo corporativo">
            <input type="password" name="nueva_password" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-sm text-white" placeholder="Nueva contraseña">
            <button type="submit" class="w-full bg-teal-700 text-white py-2.5 rounded-lg font-bold text-sm">Actualizar</button>
        </form>
        <div class="text-center mt-6"><a href="/login-form" class="text-xs text-teal-400 font-semibold hover:underline">← Volver al Login</a></div>
    </div>
</body>
</html>
"""

EDITAR_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head><meta charset="UTF-8"><title>Editar Permisos</title><script src="https://cdn.tailwindcss.com"></script></head>
<body class="bg-black text-slate-100 min-h-screen flex items-center justify-center p-6">
    <div class="bg-slate-900 border border-slate-800 p-8 rounded-2xl shadow-2xl w-full max-w-md">
        <h1 class="text-xl font-bold text-purple-400 mb-2">Editar Permisos</h1>
        <form action="/actualizar-usuario/{{ usuario.id }}" method="POST" class="space-y-4">
            <input type="text" name="nombre" value="{{ usuario.nombre }}" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-sm text-white">
            <input type="text" name="password" value="{{ usuario.password }}" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-sm text-white">
            <button type="submit" class="bg-purple-700 text-white px-5 py-2 rounded-xl text-xs font-bold w-full">Guardar Cambios</button>
        </form>
    </div>
</body>
</html>
"""

HTML_TEMPLATE_HILOS = """
<!DOCTYPE html>
<html lang="es">
<head><meta charset="UTF-8"><title>DENIER - Hilos</title><script src="https://cdn.tailwindcss.com"></script></head>
<body class="bg-black text-slate-100 min-h-screen p-6">
    <div class="max-w-7xl mx-auto">
        <div class="flex justify-between items-center mb-6 bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl">
            <div>
                <h1 class="text-2xl font-black tracking-wider uppercase text-white">DENIER - HILOS</h1>
                <p class="text-sm text-slate-400">Módulo de Control de Conos y Hilos</p>
            </div>
        </div>
        <div class="flex flex-wrap gap-3 mb-6">
            <a href="/" class="px-4 py-3 rounded-xl font-bold bg-slate-800 text-white shadow hover:bg-slate-700">← Menú DENIER</a>
            <a href="/hilos" class="px-4 py-3 rounded-xl font-semibold bg-teal-700 text-white shadow">Inventario Activo</a>
            <a href="/hilos/agotados" class="px-4 py-3 rounded-xl font-semibold bg-slate-800 text-slate-200 border border-slate-700 shadow">Hilos Sin Stock</a>
            <a href="/hilos/historial" class="px-4 py-3 rounded-xl font-semibold bg-slate-800 text-slate-200 border border-slate-700 shadow">Historial Packing List</a>
            <a href="/hilos/devoluciones-modulo" class="px-4 py-3 rounded-xl font-semibold bg-slate-800 text-slate-200 border border-slate-700 shadow">Módulo Devoluciones</a>
            <a href="/hilos/exportar-excel" class="px-4 py-3 rounded-xl font-semibold bg-emerald-800 text-white shadow ml-auto">Descargar Excel</a>
        </div>
        <!-- FORMULARIO NUEVO HILO -->
        <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl mb-8">
            <h2 class="text-xl font-bold mb-4 text-teal-400">Registrar Nuevo Hilo</h2>
            <form action="/hilos/guardar" method="POST" class="grid grid-cols-1 md:grid-cols-3 gap-4">
                <div>
                    <label class="block text-sm mb-1 text-slate-300">Tipo de Hilo</label>
                    <select name="tipo" id="tipoSelect" onchange="toggleOC()" class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white">
                        <option value="Denier">Denier (Propios)</option>
                        <option value="Artitex">Artitex (Cliente)</option>
                    </select>
                </div>
                <div id="ocDiv"><label class="block text-sm mb-1 text-slate-300">Orden / Ingreso</label><input type="text" name="oc" class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Código</label><input type="text" name="codigo" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Calibre</label><input type="text" name="calibre" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Composición</label><input type="text" name="composicion" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Proveedor</label><input type="text" name="proveedor" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Lote</label><input type="text" name="lote" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Kilos Iniciales</label><input type="number" step="0.01" name="kilos" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Conos Iniciales</label><input type="number" name="conos" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Ubicación</label><input type="text" name="ubicacion" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white" placeholder="Ej: A-1"></div>
                <div class="md:col-span-3 flex justify-end"><button type="submit" class="bg-teal-700 hover:bg-teal-600 text-white px-6 py-2.5 rounded-xl font-medium">Guardar Hilo</button></div>
            </form>
        </div>
        <!-- BUSCADOR -->
        <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl mb-8">
            <form method="GET" action="/hilos" class="grid grid-cols-1 md:grid-cols-4 gap-4">
                <select name="filtro_tipo" class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"><option value="">-- Todos Tipos --</option><option value="Denier" {{ 'selected' if filtro_tipo == 'Denier' else '' }}>Denier</option><option value="Artitex" {{ 'selected' if filtro_tipo == 'Artitex' else '' }}>Artitex</option></select>
                <select name="filtro_prov" class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"><option value="">-- Todos Proveedores --</option>{% for prov in proveedores_unicos %}<option value="{{ prov }}" {{ 'selected' if filtro_prov == prov else '' }}>{{ prov }}</option>{% endfor %}</select>
                <input type="text" name="q" value="{{ query }}" placeholder="Buscar Lote, Ubicación (A-1)..." class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white">
                <button type="submit" class="w-full bg-slate-800 text-white p-2.5 rounded-lg font-medium border border-slate-700">Filtrar y Sumar</button>
            </form>
        </div>
        <form action="/hilos/generar-packing-form" method="POST">
            <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl mb-6 flex justify-between items-center">
                <span class="text-sm text-slate-300">Suma Filtrada: <strong class="text-teal-400 text-base">{{ "%.2f"|format(total_kilos_filtrados) }} kg</strong> / <strong class="text-teal-400 text-base">{{ total_conos_filtrados }} conos</strong></span>
                <button type="submit" class="bg-purple-700 hover:bg-purple-600 text-white px-6 py-2.5 rounded-xl font-medium shadow">Crear Packing List Hilos</button>
            </div>
            <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl overflow-x-auto">
                <table class="w-full text-left">
                    <thead><tr class="border-b border-slate-800 text-sm text-slate-400"><th class="p-3">Sel</th><th class="p-3">Tipo</th><th class="p-3">Orden</th><th class="p-3">Lote</th><th class="p-3">Composición</th><th class="p-3">Kilos</th><th class="p-3">Conos</th><th class="p-3">Ubicación</th></tr></thead>
                    <tbody class="divide-y divide-slate-800 text-sm">
                        {% for item in resultados %}
                        <tr class="{{ 'bg-amber-950/40' if item.kilos < 50 else '' }}">
                            <td class="p-3 text-center"><input type="checkbox" name="hilos_sel" value="{{ item.id }}" class="w-4 h-4 text-purple-600 rounded"></td>
                            <td class="p-3">{{ item.tipo }}</td><td class="p-3 font-bold text-teal-400">{{ item.oc }}</td><td class="p-3">{{ item.lote }}</td><td class="p-3">{{ item.composicion }}</td>
                            <td class="p-3 font-bold text-teal-300">{{ item.kilos }} kg</td><td class="p-3 text-purple-400 font-bold">{{ item.conos }}</td><td class="p-3 font-semibold text-emerald-400">{{ item.ubicacion }}</td>
                        </tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </form>
    </div>
    <script>function toggleOC() { const t = document.getElementById('tipoSelect').value; document.getElementById('ocDiv').style.display = t === 'Artitex' ? 'none' : 'block'; }</script>
</body>
</html>
"""

HTML_TEMPLATE_TELA = """
<!DOCTYPE html>
<html lang="es">
<head><meta charset="UTF-8"><title>DENIER - Tela Cruda</title><script src="https://cdn.tailwindcss.com"></script></head>
<body class="bg-black text-slate-100 min-h-screen p-6">
    <div class="max-w-7xl mx-auto">
        <div class="flex justify-between items-center mb-6 bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl">
            <div>
                <h1 class="text-2xl font-black tracking-wider uppercase text-white">DENIER - TELA CRUDA</h1>
                <p class="text-sm text-slate-400">Módulo de Tejeduría y Rollos Crudos</p>
            </div>
        </div>
        <div class="flex flex-wrap gap-3 mb-6">
            <a href="/" class="px-4 py-3 rounded-xl font-bold bg-slate-800 text-white shadow hover:bg-slate-700">← Menú DENIER</a>
            <a href="/cruda" class="px-4 py-3 rounded-xl font-semibold bg-teal-700 text-white shadow">Inventario Cruda</a>
            <a href="/cruda/dashboard" class="px-4 py-3 rounded-xl font-semibold bg-slate-800 text-slate-200 border border-slate-700 shadow">Dashboard KPIs</a>
            <a href="/cruda/carga-masiva-form" class="px-4 py-3 rounded-xl font-semibold bg-slate-800 text-slate-200 border border-slate-700 shadow">Carga Masiva (Knitopia)</a>
            <a href="/cruda/historial-tela" class="px-4 py-3 rounded-xl font-semibold bg-slate-800 text-slate-200 border border-slate-700 shadow">Historial Despachos</a>
            <a href="/cruda/reingresos-modulo" class="px-4 py-3 rounded-xl font-semibold bg-slate-800 text-slate-200 border border-slate-700 shadow">Módulo Reingresos</a>
            <a href="/cruda/exportar-excel-tela" class="px-4 py-3 rounded-xl font-semibold bg-emerald-800 text-white shadow ml-auto">Descargar Excel</a>
        </div>
        <!-- FORMULARIO ROLLO -->
        <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl mb-8">
            <h2 class="text-xl font-bold mb-4 text-teal-400">Registrar Rollo Individual</h2>
            <form action="/cruda/guardar-tela" method="POST" class="grid grid-cols-1 md:grid-cols-3 gap-4">
                <div><label class="block text-sm mb-1 text-slate-300">OC</label><input type="text" name="oc" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Proveedor</label><input type="text" name="proveedor" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Lote Tela</label><input type="text" name="lote_tela" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">No. Rollo</label><input type="text" name="rollo" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Kilos</label><input type="number" step="0.01" name="kilos" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Ubicación</label><input type="text" name="ubicacion" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white" placeholder="Ej: A-1"></div>
                <div class="md:col-span-3 flex justify-end"><button type="submit" class="bg-teal-700 hover:bg-teal-600 text-white px-6 py-2.5 rounded-xl font-medium">Guardar Rollo</button></div>
            </form>
        </div>
        <!-- BUSCADOR TELA -->
        <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl mb-8">
            <form method="GET" action="/cruda" class="grid grid-cols-1 md:grid-cols-3 gap-4">
                <select name="filtro_prov" class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"><option value="">-- Proveedor --</option>{% for prov in proveedores_unicos %}<option value="{{ prov }}" {{ 'selected' if filtro_prov == prov else '' }}>{{ prov }}</option>{% endfor %}</select>
                <input type="text" name="q" value="{{ query }}" placeholder="Buscar Lote, Rollo, Ubicación (A-1)..." class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white">
                <button type="submit" class="bg-slate-800 border border-slate-700 text-white p-2.5 rounded-lg font-medium">Filtrar</button>
            </form>
        </div>
        <form action="/cruda/generar-packing-tela-form" method="POST">
            <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl mb-6 flex justify-between items-center">
                <span class="text-sm text-slate-400">Total Kilos Filtrados: <strong class="text-teal-400 text-base">{{ "%.2f"|format(total_kilos_filtrados) }} kg</strong></span>
                <button type="submit" class="bg-purple-700 hover:bg-purple-600 text-white px-6 py-2.5 rounded-xl font-medium shadow">Crear Packing List Tela</button>
            </div>
            <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl overflow-x-auto">
                <table class="w-full text-left">
                    <thead><tr class="border-b border-slate-800 text-sm text-slate-400"><th class="p-3">Sel</th><th class="p-3">OC</th><th class="p-3">Proveedor</th><th class="p-3">Lote</th><th class="p-3">Rollo</th><th class="p-3">Kilos</th><th class="p-3">Ubicación</th></tr></thead>
                    <tbody class="divide-y divide-slate-800 text-sm">
                        {% for item in resultados %}
                        <tr>
                            <td class="p-3 text-center"><input type="checkbox" name="telas_sel" value="{{ item.id }}" class="w-4 h-4 text-purple-600 rounded"></td>
                            <td class="p-3 font-bold text-teal-400">{{ item.oc }}</td><td class="p-3">{{ item.proveedor }}</td><td class="p-3">{{ item.lote_tela }}</td><td class="p-3">{{ item.rollo }}</td>
                            <td class="p-3 font-bold text-teal-300">{{ item.kilos }} kg</td><td class="p-3 font-semibold text-emerald-400">{{ item.ubicacion }}</td>
                        </tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </form>
    </div>
</body>
</html>
"""

HTML_TEMPLATE_DEVOLUCION = """
<!DOCTYPE html>
<html lang="es">
<head><meta charset="UTF-8"><title>DENIER - Devoluciones Procesada</title><script src="https://cdn.tailwindcss.com"></script></head>
<body class="bg-black text-slate-100 min-h-screen p-6">
    <div class="max-w-7xl mx-auto">
        <div class="flex justify-between items-center mb-6 bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl">
            <div>
                <h1 class="text-2xl font-black tracking-wider uppercase text-white">DENIER - TELA PROCESADA DEVUELTA</h1>
                <p class="text-sm text-slate-400">Control de devoluciones y gestión avanzada de archivos PDF/Escáner</p>
            </div>
        </div>
        <div class="flex flex-wrap gap-3 mb-6">
            <a href="/" class="px-4 py-3 rounded-xl font-bold bg-slate-800 text-white shadow hover:bg-slate-700">← Menú DENIER</a>
            <a href="/devo" class="px-4 py-3 rounded-xl font-semibold bg-rose-700 text-white shadow">Devoluciones Activas</a>
            <a href="/devo/exportar-excel-devolucion" class="px-4 py-3 rounded-xl font-semibold bg-emerald-800 text-white shadow ml-auto">Descargar Excel</a>
        </div>
        {% if mensaje %}<div class="bg-emerald-950 border border-emerald-800 text-emerald-300 p-3 rounded-xl mb-6 text-sm font-bold text-center">{{ mensaje }}</div>{% endif %}
        {% if error %}<div class="bg-red-950 border border-red-800 text-red-300 p-3 rounded-xl mb-6 text-sm font-bold text-center">{{ error }}</div>{% endif %}
        
        <!-- FORMULARIO DEVOLUCION -->
        <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl mb-8">
            <h2 class="text-xl font-bold mb-4 text-rose-400">Registrar Nueva Tela Devuelta</h2>
            <form action="/devo/guardar-devolucion" method="POST" enctype="multipart/form-data" class="grid grid-cols-1 md:grid-cols-3 gap-4">
                <div><label class="block text-sm mb-1 text-slate-300">Cliente</label><input type="text" name="cliente" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Lote</label><input type="text" name="lote" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Tela</label><input type="text" name="tela" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Color</label><input type="text" name="color" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Rollos</label><input type="number" name="rollos" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Kilos</label><input type="number" step="0.01" name="kilos" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Yardas</label><input type="number" step="0.01" name="yardas" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"></div>
                <div><label class="block text-sm mb-1 text-slate-300">Calidad</label><select name="calidad" class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white"><option value="Primera">Primera</option><option value="Segunda">Segunda</option><option value="Rechazada">Rechazada</option></select></div>
                <div><label class="block text-sm mb-1 text-slate-300">Ubicación</label><input type="text" name="ubicacion" required class="w-full p-2.5 bg-black border border-slate-700 rounded-lg text-white" placeholder="Rack 3 / A-1"></div>
                <div class="md:col-span-3"><label class="block text-sm mb-1 text-slate-300">Archivo de Respaldo (PDF/Imagen)</label><input type="file" name="archivo_devolucion" accept=".pdf,.png,.jpg,.jpeg" class="w-full p-2 bg-black border border-slate-700 rounded-lg text-sm text-slate-300"></div>
                <div class="md:col-span-3 flex justify-end"><button type="submit" class="bg-rose-700 hover:bg-rose-600 text-white px-6 py-2.5 rounded-xl font-medium">Registrar Devolución</button></div>
            </form>
        </div>
        <!-- TABLA DEVOLUCIONES -->
        <div class="bg-slate-900 border border-slate-800 p-6 rounded-2xl shadow-xl overflow-x-auto">
            <table class="w-full text-left">
                <thead><tr class="border-b border-slate-800 text-sm text-slate-400"><th class="p-3">Cliente</th><th class="p-3">Lote</th><th class="p-3">Tela / Color</th><th class="p-3">Rollos</th><th class="p-3">Kilos / Yardas</th><th class="p-3">Calidad / Ubicación</th><th class="p-3">Archivo PDF</th></tr></thead>
                <tbody class="divide-y divide-slate-800 text-sm">
                    {% for item in resultados %}
                    <tr>
                        <td class="p-3 font-bold text-rose-400">{{ item.cliente }}</td><td class="p-3">{{ item.lote }}</td><td class="p-3">{{ item.tela }}<br><span class="text-xs text-slate-400">{{ item.color }}</span></td>
                        <td class="p-3 font-bold">{{ item.rollos }} rollos</td><td class="p-3 text-rose-300 font-bold">{{ "%.2f"|format(item.kilos) }} kg<br><span class="text-xs text-slate-400">{{ "%.2f"|format(item.yardas) }} yds</span></td>
                        <td><span class="px-2 py-0.5 rounded text-xs font-bold bg-rose-950 text-rose-400 border border-rose-800">{{ item.calidad }}</span><div class="text-xs text-emerald-400 mt-1 font-semibold">{{ item.ubicacion }}</div></td>
                        <td class="p-3">
                            {% if item.archivo %}
                                <div class="mb-2 bg-slate-950 p-2 rounded border border-slate-800 flex items-center justify-between">
                                    <a href="/devo/ver-archivo/{{ item.archivo }}" target="_blank" class="text-teal-400 text-xs">📄 Ver Archivo</a>
                                    <a href="/devo/eliminar-archivo/{{ item.id }}" class="text-red-400 text-xs bg-red-950 border border-red-800 px-1 rounded">Eliminar</a>
                                </div>
                            {% else %}
                                <span class="text-xs text-amber-400">Sin archivo</span>
                            {% endif %}
                            <form action="/devo/subir-archivo-devolucion/{{ item.id }}" method="POST" enctype="multipart/form-data" class="flex gap-2 items-center mt-1">
                                <input type="file" name="archivo_devolucion" accept=".pdf,.png,.jpg,.jpeg" required class="text-xs text-slate-400">
                                <button type="submit" class="bg-slate-800 text-white px-2 py-1 rounded text-xs font-bold">Subir</button>
                            </form>
                        </td>
                    </tr>
                    {% endfor %}
                </tbody>
            </table>
        </div>
    </div>
</body>
</html>
"""

# ==========================================
# RUTAS PRINCIPALES Y LOGIN
# ==========================================
@app.route('/')
def index():
    if 'email' not in session:
        return redirect(url_for('login_form'))
    return render_template_string(HOME_TEMPLATE, usuarios_db=usuarios_db)

@app.route('/login-form')
def login_form():
    return render_template_string(LOGIN_TEMPLATE, error=None)

@app.route('/login', methods=['POST'])
def login():
    email = request.form.get('email', '').strip()
    password = request.form.get('password', '').strip()
    usuario = next((u for u in usuarios_db if u['email'] == email and u['password'] == password), None)
    if usuario:
        session['email'] = usuario['email']
        session['nombre'] = usuario['nombre']
        session['modulos'] = usuario['modulos']
        return redirect(url_for('index'))
    return render_template_string(LOGIN_TEMPLATE, error="Correo o contraseña incorrectos")

@app.route('/reset-password-form')
def reset_password_form():
    return render_template_string(RESET_TEMPLATE, error=None, mensaje=None)

@app.route('/procesar-reset', methods=['POST'])
def procesar_reset():
    email = request.form.get('email', '').strip()
    nueva = request.form.get('nueva_password', '').strip()
    usuario = next((u for u in usuarios_db if u['email'] == email), None)
    if not usuario:
        return render_template_string(RESET_TEMPLATE, error="El correo ingresado no está registrado.", mensaje=None)
    usuario['password'] = nueva
    return render_template_string(RESET_TEMPLATE, error=None, mensaje="¡Contraseña actualizada con éxito!")

@app.route('/crear-usuario', methods=['POST'])
def crear_usuario():
    if 'admin' not in session.get('modulos', []):
        return redirect(url_for('index'))
    nuevo_email = request.form.get('email', '').strip()
    if any(u['email'] == nuevo_email for u in usuarios_db):
        return redirect(url_for('index'))
    nuevo = {
        'id': max([u['id'] for u in usuarios_db], default=0) + 1,
        'email': nuevo_email,
        'password': request.form.get('password', '').strip(),
        'nombre': request.form.get('nombre', '').strip(),
        'modulos': request.form.getlist('modulos')
    }
    usuarios_db.append(nuevo)
    return redirect(url_for('index'))

@app.route('/editar-usuario/<int:uid>')
def editar_usuario(uid):
    if 'admin' not in session.get('modulos', []):
        return redirect(url_for('index'))
    usuario = next((u for u in usuarios_db if u['id'] == uid), None)
    return render_template_string(EDITAR_TEMPLATE, usuario=usuario)

@app.route('/actualizar-usuario/<int:uid>', methods=['POST'])
def actualizar_usuario(uid):
    if 'admin' not in session.get('modulos', []):
        return redirect(url_for('index'))
    usuario = next((u for u in usuarios_db if u['id'] == uid), None)
    if usuario:
        usuario['nombre'] = request.form.get('nombre', '').strip()
        usuario['password'] = request.form.get('password', '').strip()
    return redirect(url_for('index'))

@app.route('/eliminar-usuario/<int:uid>')
def eliminar_usuario(uid):
    global usuarios_db
    usuarios_db = [u for u in usuarios_db if u['id'] != uid or u['email'] == 'controlprod_1@deniersa.com']
    return redirect(url_for('index'))

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login_form'))

# ==========================================
# RUTAS MÓDULO HILOS (/hilos)
# ==========================================
@app.route('/hilos')
def hilos_index():
    if 'email' not in session or ('hilos' not in session.get('modulos', []) and 'admin' not in session.get('modulos', [])):
        return redirect(url_for('index'))
    query = request.args.get('q', '').lower()
    filtro_tipo = request.args.get('filtro_tipo', '')
    filtro_prov = request.args.get('filtro_prov', '')
    proveedores_unicos = sorted(list(set(i['proveedor'] for i in inventario_hilos)))
    activos = [i for i in inventario_hilos if i['kilos'] > 0 or i['conos'] > 0]
    resultados = activos
    if filtro_tipo: resultados = [i for i in resultados if i['tipo'] == filtro_tipo]
    if filtro_prov: resultados = [i for i in resultados if i['proveedor'] == filtro_prov]
    if query: resultados = [i for i in resultados if query in i['lote'].lower() or query in i['composicion'].lower() or query in i['oc'].lower() or query in i['ubicacion'].lower()]
    return render_template_string(HTML_TEMPLATE_HILOS, resultados=resultados, proveedores_unicos=proveedores_unicos, filtro_tipo=filtro_tipo, filtro_prov=filtro_prov, total_kilos_filtrados=sum(i['kilos'] for i in resultados), total_conos_filtrados=sum(i['conos'] for i in resultados), query=query)

@app.route('/hilos/guardar', methods=['POST'])
def hilos_guardar():
    tipo = request.form.get('tipo')
    inventario_hilos.append({
        'id': len(inventario_hilos) + 1,
        'tipo': tipo,
        'oc': 'N/A' if tipo == 'Artitex' else request.form.get('oc'),
        'codigo': request.form.get('codigo'),
        'calibre': request.form.get('calibre'),
        'composicion': request.form.get('composicion'),
        'proveedor': request.form.get('proveedor'),
        'lote': request.form.get('lote'),
        'kilos': float(request.form.get('kilos')),
        'conos': int(request.form.get('conos')),
        'ubicacion': request.form.get('ubicacion')
    })
    return redirect(url_for('hilos_index'))

@app.route('/hilos/exportar-excel')
def hilos_exportar_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tipo", "OC", "Código", "Calibre", "Composición", "Proveedor", "Lote", "Kilos", "Conos", "Ubicación"])
    for i in inventario_hilos: ws.append([i['tipo'], i['oc'], i['codigo'], i['calibre'], i['composicion'], i['proveedor'], i['lote'], i['kilos'], i['conos'], i['ubicacion']])
    io_b = io.BytesIO()
    wb.save(io_b)
    io_b.seek(0)
    return send_file(io_b, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='Hilos.xlsx')

@app.route('/hilos/agotados')
def hilos_agotados():
    return render_template_string("""<!DOCTYPE html><html lang="es"><head><script src="https://cdn.tailwindcss.com"></script></head><body class="bg-black text-slate-100 p-8"><div class="max-w-4xl mx-auto bg-slate-900 p-6 rounded-xl"><h1 class="text-xl font-bold mb-4">Hilos sin stock</h1><a href="/hilos" class="bg-slate-800 text-white px-4 py-2 rounded">Volver</a></div></body></html>""")

@app.route('/hilos/historial')
def hilos_historial():
    return render_template_string("""<!DOCTYPE html><html lang="es"><head><script src="https://cdn.tailwindcss.com"></script></head><body class="bg-black text-slate-100 p-8"><div class="max-w-4xl mx-auto bg-slate-900 p-6 rounded-xl"><h1 class="text-xl font-bold mb-4">Historial Hilos</h1><a href="/hilos" class="bg-slate-800 text-white px-4 py-2 rounded">Volver</a></div></body></html>""")

@app.route('/hilos/devoluciones-modulo')
def hilos_devoluciones_modulo():
    return render_template_string("""<!DOCTYPE html><html lang="es"><head><script src="https://cdn.tailwindcss.com"></script></head><body class="bg-black text-slate-100 p-8"><div class="max-w-4xl mx-auto bg-slate-900 p-6 rounded-xl"><h1 class="text-xl font-bold mb-4">Devoluciones Hilos</h1><a href="/hilos" class="bg-slate-800 text-white px-4 py-2 rounded">Volver</a></div></body></html>""")

@app.route('/hilos/generar-packing-form', methods=['POST'])
def hilos_generar_packing_form():
    return redirect(url_for('hilos_index'))

# ==========================================
# RUTAS MÓDULO TELA CRUDA (/cruda)
# ==========================================
@app.route('/cruda')
def cruda_index():
    if 'email' not in session or ('cruda' not in session.get('modulos', []) and 'admin' not in session.get('modulos', [])):
        return redirect(url_for('index'))
    query = request.args.get('q', '').lower()
    filtro_prov = request.args.get('filtro_prov', '')
    proveedores_unicos = sorted(list(set(i['proveedor'] for i in inventario_tela)))
    activos = [i for i in inventario_tela if i['kilos'] > 0]
    resultados = activos
    if filtro_prov: resultados = [i for i in resultados if i['proveedor'] == filtro_prov]
    if query: resultados = [i for i in resultados if query in i['lote_tela'].lower() or query in i['oc'].lower() or query in i['rollo'].lower() or query in i['proveedor'].lower() or query in i['ubicacion'].lower()]
    return render_template_string(HTML_TEMPLATE_TELA, resultados=resultados, proveedores_unicos=proveedores_unicos, filtro_prov=filtro_prov, query=query, total_kilos_filtrados=sum(i['kilos'] for i in resultados))

@app.route('/cruda/guardar-tela', methods=['POST'])
def cruda_guardar_tela():
    inventario_tela.append({
        'id': len(inventario_tela) + 1,
        'oc': request.form.get('oc'),
        'proveedor': request.form.get('proveedor'),
        'lote_tela': request.form.get('lote_tela'),
        'rollo': request.form.get('rollo'),
        'kilos': float(request.form.get('kilos')),
        'ubicacion': request.form.get('ubicacion')
    })
    return redirect(url_for('cruda_index'))

@app.route('/cruda/exportar-excel-tela')
def cruda_exportar_excel_tela():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["OC", "Proveedor", "Lote", "Rollo", "Kilos", "Ubicación"])
    for i in inventario_tela: ws.append([i['oc'], i['proveedor'], i['lote_tela'], i['rollo'], i['kilos'], i['ubicacion']])
    io_b = io.BytesIO()
    wb.save(io_b)
    io_b.seek(0)
    return send_file(io_b, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='Tela.xlsx')

@app.route('/cruda/dashboard')
def cruda_dashboard():
    return render_template_string("""<!DOCTYPE html><html lang="es"><head><script src="https://cdn.tailwindcss.com"></script></head><body class="bg-black text-slate-100 p-8"><div class="max-w-4xl mx-auto bg-slate-900 p-6 rounded-xl"><h1 class="text-xl font-bold mb-4">Dashboard KPIs</h1><a href="/cruda" class="bg-slate-800 text-white px-4 py-2 rounded">Volver</a></div></body></html>""")

@app.route('/cruda/carga-masiva-form')
def cruda_carga_masiva_form():
    return render_template_string("""<!DOCTYPE html><html lang="es"><head><script src="https://cdn.tailwindcss.com"></script></head><body class="bg-black text-slate-100 p-8"><div class="max-w-4xl mx-auto bg-slate-900 p-6 rounded-xl"><h1 class="text-xl font-bold mb-4">Carga Masiva</h1><a href="/cruda" class="bg-slate-800 text-white px-4 py-2 rounded">Volver</a></div></body></html>""")

@app.route('/cruda/historial-tela')
def cruda_historial_tela():
    return render_template_string("""<!DOCTYPE html><html lang="es"><head><script src="https://cdn.tailwindcss.com"></script></head><body class="bg-black text-slate-100 p-8"><div class="max-w-4xl mx-auto bg-slate-900 p-6 rounded-xl"><h1 class="text-xl font-bold mb-4">Historial Tela</h1><a href="/cruda" class="bg-slate-800 text-white px-4 py-2 rounded">Volver</a></div></body></html>""")

@app.route('/cruda/reingresos-modulo')
def cruda_reingresos_modulo():
    return render_template_string("""<!DOCTYPE html><html lang="es"><head><script src="https://cdn.tailwindcss.com"></script></head><body class="bg-black text-slate-100 p-8"><div class="max-w-4xl mx-auto bg-slate-900 p-6 rounded-xl"><h1 class="text-xl font-bold mb-4">Reingresos Tela</h1><a href="/cruda" class="bg-slate-800 text-white px-4 py-2 rounded">Volver</a></div></body></html>""")

@app.route('/cruda/generar-packing-tela-form', methods=['POST'])
def cruda_generar_packing_tela_form():
    return redirect(url_for('cruda_index'))

# ==========================================
# RUTAS MÓDULO DEVOLUCIONES (/devo)
# ==========================================
@app.route('/devo')
def devo_index():
    if 'email' not in session or ('devo' not in session.get('modulos', []) and 'admin' not in session.get('modulos', [])):
        return redirect(url_for('index'))
    query = request.args.get('q', '').lower()
    resultados = devoluciones_procesada
    if query: resultados = [i for i in resultados if query in i['cliente'].lower() or query in i['lote'].lower() or query in i['tela'].lower()]
    return render_template_string(HTML_TEMPLATE_DEVOLUCION, resultados=resultados, query=query, mensaje=request.args.get('mensaje',''), error=request.args.get('error',''))

@app.route('/devo/guardar-devolucion', methods=['POST'])
def devo_guardar_devolucion():
    archivo = request.files.get('archivo_devolucion')
    nombre_archivo = None
    if archivo and archivo.filename:
        nombre_archivo = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{archivo.filename}"
        archivo.save(os.path.join(app.config['UPLOAD_FOLDER'], nombre_archivo))
    devoluciones_procesada.append({
        'id': len(devoluciones_procesada) + 1,
        'cliente': request.form.get('cliente', '').strip(),
        'lote': request.form.get('lote', '').strip(),
        'tela': request.form.get('tela', '').strip(),
        'color': request.form.get('color', '').strip(),
        'rollos': int(request.form.get('rollos', 0)),
        'kilos': float(request.form.get('kilos', 0)),
        'yardas': float(request.form.get('yardas', 0)),
        'calidad': request.form.get('calidad'),
        'ubicacion': request.form.get('ubicacion', '').strip(),
        'archivo': nombre_archivo
    })
    return redirect(url_for('devo_index', mensaje="¡Devolución registrada con éxito!"))

@app.route('/devo/subir-archivo-devolucion/<int:dev_id>', methods=['POST'])
def devo_subir_archivo_devolucion(dev_id):
    archivo = request.files.get('archivo_devolucion')
    if archivo and archivo.filename:
        item = next((i for i in devoluciones_procesada if i['id'] == dev_id), None)
        if item:
            nombre_archivo = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{archivo.filename}"
            archivo.save(os.path.join(app.config['UPLOAD_FOLDER'], nombre_archivo))
            item['archivo'] = nombre_archivo
            return redirect(url_for('devo_index', mensaje="¡Archivo actualizado con éxito!"))
    return redirect(url_for('devo_index', error="No se seleccionó archivo válido."))

@app.route('/devo/ver-archivo/<filename>')
def devo_ver_archivo(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

@app.route('/devo/eliminar-archivo/<int:dev_id>')
def devo_eliminar_archivo(dev_id):
    item = next((i for i in devoluciones_procesada if i['id'] == dev_id), None)
    if item: item['archivo'] = None
    return redirect(url_for('devo_index', mensaje="¡Archivo eliminado!"))

@app.route('/devo/exportar-excel-devolucion')
def devo_exportar_excel_devolucion():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Cliente", "Lote", "Tela", "Color", "Rollos", "Kilos", "Yardas", "Calidad", "Ubicación", "Archivo"])
    for i in devoluciones_procesada: ws.append([i['cliente'], i['lote'], i['tela'], i['color'], i['rollos'], i['kilos'], i['yardas'], i['calidad'], i['ubicacion'], i['archivo'] or 'Sin archivo'])
    io_b = io.BytesIO()
    wb.save(io_b)
    io_b.seek(0)
    return send_file(io_b, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='Devoluciones.xlsx')

if __name__ == '__main__':
    app.run(debug=True, port=5000)
